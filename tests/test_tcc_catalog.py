import unittest
from datetime import datetime, timezone, timedelta
from pathlib import Path
import tempfile
import openpyxl

from src.utils.paths import get_output_dir, get_medias_dir, get_sheets_dir
from src.pipeline.cleaner import anonymize_user, extract_urls, clean_text_for_nlp
from src.pipeline.filters import parse_date_to_brt, BRT
from src.collectors.telegram_collector import format_reactions
from src.db.session import SessionLocal, init_db
from src.db.models import Message, ContentAnalysis
from src.exporter.excel_exporter import export_to_tcc_spreadsheet

class TestTCCCatalog(unittest.TestCase):

    def test_paths_creation(self):
        """Verifica se os diretórios output e subpastas são criados corretamente."""
        output_dir = get_output_dir()
        medias_dir = get_medias_dir()
        sheets_dir = get_sheets_dir()

        self.assertTrue(output_dir.exists())
        self.assertTrue(medias_dir.exists())
        self.assertTrue(sheets_dir.exists())
        self.assertEqual(medias_dir.name, "Mídias")
        self.assertEqual(sheets_dir.name, "Planilhas de Catalogação")

    def test_media_filename_standard(self):
        """Verifica a padronização: Iddamensagem_data_horário.ext."""
        msg_id = 1542
        dt = datetime(2026, 9, 15, 14, 30, 45)
        date_str = dt.strftime("%Y-%m-%d")
        time_str = dt.strftime("%H-%M-%S")
        filename = f"{msg_id}_{date_str}_{time_str}.jpg"

        self.assertEqual(filename, "1542_2026-09-15_14-30-45.jpg")

    def test_strict_anonymization(self):
        """Verifica conformidade com Ética em Pesquisa: nenhum dado pessoal deve vazar."""
        raw_user_id = 987654321
        civilian_name = "João da Silva Ferreira"
        
        anon_id, anon_name = anonymize_user(raw_user_id, civilian_name)
        
        self.assertTrue(anon_id.startswith("User_"))
        self.assertNotIn("João", anon_id)
        self.assertNotIn("Silva", anon_id)
        self.assertNotIn("Ferreira", anon_id)
        self.assertNotIn(str(raw_user_id), anon_id)

    def test_date_conversion_brt(self):
        """Verifica se a conversão UTC -> BRT (UTC-3) funciona com precisão."""
        dt_utc = datetime(2026, 9, 14, 15, 0, 0, tzinfo=timezone.utc)
        dt_brt = parse_date_to_brt(dt_utc)
        
        self.assertEqual(dt_brt.hour, 12)
        self.assertEqual(dt_brt.day, 14)

    def test_reactions_formatting(self):
        """Verifica a formatação de reações do Telegram."""
        class MockReaction:
            def __init__(self, emoticon):
                self.emoticon = emoticon

        class MockResult:
            def __init__(self, emoticon, count):
                self.reaction = MockReaction(emoticon)
                self.count = count

        class MockReactionsObj:
            def __init__(self, results):
                self.results = results

        mock_reactions = MockReactionsObj([
            MockResult("👍", 10),
            MockResult("❤️", 5)
        ])

        total, summary, r_dict = format_reactions(mock_reactions)
        self.assertEqual(total, 15)
        self.assertIn("👍 (10)", summary)
        self.assertIn("❤️ (5)", summary)
        self.assertEqual(r_dict["👍"], 10)
        self.assertEqual(r_dict["❤️"], 5)

    def test_excel_export_structure_and_anonymity(self):
        """Verifica a criação da planilha com colunas corretas e ausência de dados pessoais."""
        init_db()
        db = SessionLocal()

        # Insere uma mensagem de teste
        test_msg = Message(
            telegram_msg_id=999999,
            chat_id="test_chat",
            chat_title="Grupo de Teste",
            date_utc=datetime.now(timezone.utc),
            date_brt=datetime(2026, 9, 16, 10, 0, 0),
            week_label="Semana 1 (14 a 20 set)",
            sender_id_anon="User_TESTE1",
            sender_type="user",
            is_bot=False,
            text_raw="Texto de mensagem de teste para o TCC",
            text_clean="Texto de mensagem de teste para o TCC",
            media_type="Foto",
            has_media=True,
            media_filename="999999_2026-09-16_10-00-00.jpg",
            is_forward=False,
            views_count=42,
            reactions_count=3
        )
        test_msg.analysis = ContentAnalysis(status_validacao="Pendente")
        
        db.query(ContentAnalysis).delete()
        db.query(Message).filter_by(telegram_msg_id=999999).delete()
        db.commit()

        db.add(test_msg)
        db.commit()
        db.close()

        # Exporta
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp:
            tmp_path = Path(tmp.name)

        try:
            exported_path = export_to_tcc_spreadsheet(chat_id="test_chat", output_file=tmp_path)
            self.assertTrue(exported_path.exists())

            # Valida conteúdo do Excel via openpyxl
            wb = openpyxl.load_workbook(exported_path)
            ws = wb.active

            headers = [cell.value for cell in ws[1]]
            
            # Verifica colunas obrigatórias
            self.assertIn("ID da Mensagem", headers)
            self.assertIn("Data e Horário (BRT)", headers)
            self.assertIn("Publicada por Bot?", headers)
            self.assertIn("Fruto de Encaminhamento?", headers)
            self.assertIn("Possui Mídia?", headers)
            self.assertIn("Arquivo de Mídia Salvo", headers)
            self.assertIn("Visualizações", headers)
            self.assertIn("Total de Reações", headers)
            self.assertIn("Texto da Mensagem", headers)

            # Verifica total exato de colunas (15 colunas objetivas)
            self.assertEqual(len(headers), 15)
            self.assertIn("Publicada por Administrador?", headers)
            self.assertEqual(headers.index("Publicada por Administrador?"), headers.index("Publicada por Bot?") + 1)

            # Verifica que colunas manuais subjetivas foram removidas
            self.assertNotIn("Enquadramento Noticioso", headers)
            self.assertNotIn("Valência Política", headers)
            self.assertNotIn("Apelo de Credibilidade", headers)
            self.assertNotIn("Notas Netnográficas", headers)
            self.assertNotIn("Tema Predominante", headers)
            self.assertNotIn("Status da Catalogação", headers)

            # VERIFICAÇÃO ÉTICA CRÍTICA: Nenhum cabeçalho de nome de usuário real
            for h in headers:
                self.assertNotIn("Nome do Usuário", str(h))
                self.assertNotIn("Username", str(h))
                self.assertNotIn("Telefone", str(h))

            # Verifica linha 2
            row2 = [cell.value for cell in ws[2]]
            self.assertIn(999999, row2)
            self.assertIn("999999_2026-09-16_10-00-00.jpg", row2)
            wb.close()
        finally:
            if tmp_path.exists():
                tmp_path.unlink()

if __name__ == "__main__":
    unittest.main()
